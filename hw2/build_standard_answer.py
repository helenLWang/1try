#!/usr/bin/env python3
"""Build 95-703 C HW#2 standard-answer PowerPoint from the Painter ERD template."""

from copy import deepcopy
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu
from pptx.oxml import parse_xml
import os
import shutil

BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Times New Roman"

TEMPLATE = "/home/ubuntu/.cursor/projects/workspace/uploads/Template_ERD_-_Painter_Model__1__573e.pptx"
OUT_PPTX = "/workspace/hw2/HW2_Standard_Answer.pptx"
OUT_ROOT = "/workspace/HW2_Standard_Answer.pptx"
OUT_DESKTOP = os.path.expanduser("~/Desktop/HW2_Standard_Answer.pptx")


def set_run_font(run, size=12, bold=False, underline=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    run.font.color.rgb = BLACK
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", FONT)


def _tf_setup(shape, anchor="ctr", wrap=True, l=0.04, t=0.02, r=0.04, b=0.02):
    tf = shape.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    tf.margin_left = Inches(l)
    tf.margin_right = Inches(r)
    tf.margin_top = Inches(t)
    tf.margin_bottom = Inches(b)
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    if bodyPr is not None:
        bodyPr.set("anchor", anchor)
    return tf


def _fill_line(shape, fill=WHITE, line=BLACK, line_pt=1.0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_pt)


def add_textbox(slide, x, y, w, h, text, size=16, bold=False, align=PP_ALIGN.LEFT, italic=False):
    sp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = _tf_setup(sp, anchor="t", wrap=True)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, italic=italic)
    return sp


def add_rect(slide, x, y, w, h, lines, size=14, bold=False, line_pt=1.0):
    """Entity rectangle — same language as the Painter template."""
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill_line(sp, line_pt=line_pt)
    tf = _tf_setup(sp, anchor="ctr", wrap=True, l=0.05, t=0.02, r=0.05, b=0.02)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size=size, bold=bold)
    return (x, y, w, h)


def add_diamond(slide, x, y, w, h, lines, size=13, line_pt=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill_line(sp, line_pt=line_pt)
    tf = _tf_setup(sp, anchor="ctr", wrap=False, l=0.14, r=0.14, t=0.02, b=0.02)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size=size, bold=False)
    return (x, y, w, h)


def add_oval(slide, x, y, w, h, text, size=14):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill_line(sp, line_pt=1.0)
    tf = _tf_setup(sp, anchor="ctr")
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=False)
    return (x, y, w, h)


CARD_W = 0.50
CARD_H = 0.22


def add_card(slide, x, y, text, w=CARD_W, h=CARD_H, size=11):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = WHITE
    sp.line.fill.background()
    tf = _tf_setup(sp, anchor="ctr", wrap=False, l=0.0, t=0.0, r=0.0, b=0.0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size)
    return (x, y, w, h)


def center(box):
    x, y, w, h = box
    return (x + w / 2.0, y + h / 2.0)


def _overlap(a, b, pad=0.06):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return not (ax + aw + pad <= bx or bx + bw + pad <= ax or ay + ah + pad <= by or by + bh + pad <= ay)


def _exit_point(box, toward):
    cx, cy = center(box)
    x, y, w, h = box
    tx, ty = toward
    dx, dy = tx - cx, ty - cy
    if dx == 0 and dy == 0:
        return cx, cy
    t = min(
        ((w / 2) / abs(dx)) if dx else 1e9,
        ((h / 2) / abs(dy)) if dy else 1e9,
    )
    # step a little past the border so the label is clearly outside the box
    t += 0.08 / max((dx * dx + dy * dy) ** 0.5, 0.01)
    return cx + dx * t, cy + dy * t


def _gap_anchor(box_a, box_b, toward_a=True, side=1):
    """Center of a cardinality label sitting in the gap between two shapes."""
    ca, cb = center(box_a), center(box_b)
    p1 = _exit_point(box_a, cb)
    p2 = _exit_point(box_b, ca)
    t = 0.28 if toward_a else 0.72
    mx = p1[0] + (p2[0] - p1[0]) * t
    my = p1[1] + (p2[1] - p1[1]) * t
    dx, dy = p2[0] - p1[0], p2[1] - p1[1]
    length = max((dx * dx + dy * dy) ** 0.5, 0.001)
    # perpendicular offset keeps the number off the connector and off the boxes
    mx += (-dy / length) * (0.16 * side)
    my += (dx / length) * (0.16 * side)
    return mx - CARD_W / 2, my - CARD_H / 2


def add_card_in_gap_xy(slide, x, y, text, obstacles):
    nudges = [(0, 0), (0, -0.16), (0, 0.16), (-0.16, 0), (0.16, 0),
              (0, -0.28), (0, 0.28), (-0.28, 0), (0.28, 0),
              (-0.22, -0.22), (0.22, -0.22), (-0.22, 0.22), (0.22, 0.22)]
    for dx, dy in nudges:
        nx, ny = x + dx, y + dy
        nx = min(max(nx, 0.04), 13.33 - CARD_W - 0.04)
        ny = min(max(ny, 0.48), 7.05 - CARD_H)
        cand = (nx, ny, CARD_W, CARD_H)
        if any(_overlap(cand, ob, pad=0.05) for ob in obstacles):
            continue
        box = add_card(slide, nx, ny, text)
        obstacles.append(box)
        return box
    box = add_card(slide, x, y, text)
    obstacles.append(box)
    return box


def add_line(slide, a, b, width=0.75):
    x1, y1 = a if len(a) == 2 else center(a)
    x2, y2 = b if len(b) == 2 else center(b)
    sp = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    sp.line.color.rgb = BLACK
    sp.line.width = Pt(width)
    return sp


def clear_slide(slide):
    spTree = slide.shapes._spTree
    for el in list(spTree):
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
            continue
        spTree.remove(el)


def delete_slide(prs, index):
    rId = prs.slides._sldIdLst[index].get(qn("r:id"))
    prs.part.drop_rel(rId)
    sldId = prs.slides._sldIdLst[index]
    prs.slides._sldIdLst.remove(sldId)


def add_blank_slide(prs):
    # Painter template: layout 6 = Blank
    return prs.slides.add_slide(prs.slide_layouts[6])


def build_er(slide):
    add_textbox(
        slide, 0.25, 0.10, 12.8, 0.34,
        "95–703 C  Homework #2   —   Standard Answer:  ER Diagram  (no attributes)",
        size=18, bold=False,
    )

    # Wider gaps so (min,max) labels sit on the connectors, never on the boxes.
    E = {
        "SUPPLIER":     (0.16, 2.20, 1.32, 0.36),
        "PROCUREMENT":  (2.14, 2.16, 1.48, 0.44),
        "ITEM":         (4.30, 2.20, 1.16, 0.36),
        "PASM":         (4.12, 3.04, 1.52, 0.46),
        "PRODUCT":      (4.18, 3.96, 1.40, 0.36),
        "INCLUDE":      (4.04, 4.80, 1.68, 0.54),
        "PLINE":        (4.14, 5.66, 1.48, 0.36),
        "CUSTOMER":     (7.85, 0.84, 1.40, 0.36),
        "SUBMIT":       (7.70, 1.38, 1.70, 0.52),
        "ORDER":        (7.85, 2.10, 1.40, 0.36),
        "PROCESS":      (9.90, 1.96, 1.60, 0.56),
        "SREP":         (12.00, 2.10, 1.22, 0.36),
        "OLINE":        (6.05, 2.50, 1.36, 0.40),
        "PRODUCE":      (6.28, 3.82, 1.52, 0.54),
        "WC":           (8.30, 3.96, 1.42, 0.36),
        "ASSIGN":       (10.30, 3.92, 1.36, 0.42),
        "EMP":          (12.00, 3.96, 1.18, 0.36),
        "MANAGE":       (8.30, 4.80, 1.42, 0.54),
        "MGR":          (8.38, 5.66, 1.28, 0.36),
        "ISA":          (12.72, 4.72, 0.42, 0.42),
        "ADMIN":        (10.48, 5.66, 1.68, 0.36),
    }

    pairs = [
        "SUPPLIER", "PROCUREMENT", "PROCUREMENT", "ITEM",
        "ITEM", "PASM", "PASM", "PRODUCT",
        "PRODUCT", "INCLUDE", "INCLUDE", "PLINE",
        "CUSTOMER", "SUBMIT", "SUBMIT", "ORDER",
        "ORDER", "PROCESS", "PROCESS", "SREP",
        "ORDER", "OLINE", "OLINE", "PRODUCT",
        "PRODUCT", "PRODUCE", "PRODUCE", "WC",
        "WC", "ASSIGN", "ASSIGN", "EMP",
        "WC", "MANAGE", "MANAGE", "MGR",
        "EMP", "ISA", "ISA", "SREP", "ISA", "MGR", "ISA", "ADMIN",
    ]
    for a, b in zip(pairs[0::2], pairs[1::2]):
        add_line(slide, E[a], E[b], width=0.75)

    # regular entities (1 pt, like Painter)
    add_rect(slide, *E["SUPPLIER"], ["SUPPLIER"], size=13)
    add_rect(slide, *E["ITEM"], ["ITEM"], size=13)
    add_rect(slide, *E["PRODUCT"], ["PRODUCT"], size=13)
    add_rect(slide, *E["PLINE"], ["PRODUCT_LINE"], size=13)
    add_rect(slide, *E["CUSTOMER"], ["CUSTOMER"], size=13)
    add_rect(slide, *E["ORDER"], ["ORDER"], size=13)
    add_rect(slide, *E["SREP"], ["SALES_REP"], size=13)
    add_rect(slide, *E["WC"], ["WORK_CENTER"], size=12)
    add_rect(slide, *E["EMP"], ["EMPLOYEE"], size=13)
    add_rect(slide, *E["MGR"], ["MANAGER"], size=13)
    add_rect(slide, *E["ADMIN"], ["ADMIN_ASSISTANT"], size=12)

    # composite entities (thicker border = M:N transformed)
    add_rect(slide, *E["PROCUREMENT"], ["PROCUREMENT"], size=12, line_pt=2.25)
    add_rect(slide, *E["PASM"], ["PRODUCT", "ASSEMBLY"], size=12, line_pt=2.25)
    add_rect(slide, *E["OLINE"], ["ORDER_LINE"], size=12, line_pt=2.25)
    add_rect(slide, *E["ASSIGN"], ["ASSIGNMENT"], size=12, line_pt=2.25)

    add_diamond(slide, *E["INCLUDE"], ["INCLUDE"], size=13)
    add_diamond(slide, *E["SUBMIT"], ["SUBMIT"], size=13)
    add_diamond(slide, *E["PROCESS"], ["PROCESS"], size=13)
    add_diamond(slide, *E["PRODUCE"], ["PRODUCE"], size=13)
    add_diamond(slide, *E["MANAGE"], ["MANAGE"], size=13)
    add_oval(slide, *E["ISA"], "O", size=14)

    obstacles = list(E.values())

    # Hand-placed in connector whitespace — never on a rectangle/diamond.
    labels = [
        # SUPPLIER — PROCUREMENT — ITEM
        (1.52, 1.88, "(1,1)"),
        (1.52, 2.64, "(0,N)"),
        (3.66, 1.88, "(0,N)"),
        (3.66, 2.64, "(1,1)"),
        # ITEM — PRODUCT_ASSEMBLY — PRODUCT  (left of the stack)
        (3.52, 2.20, "(1,1)"),
        (3.52, 3.10, "(1,N)"),
        (3.52, 3.52, "(1,N)"),
        (3.52, 3.98, "(1,1)"),
        # PRODUCT — INCLUDE — PRODUCT_LINE
        (5.78, 4.40, "(1,N)"),
        (5.78, 5.36, "(1,1)"),
        # CUSTOMER — SUBMIT — ORDER
        (9.32, 0.86, "(1,1)"),
        (7.28, 2.10, "(0,N)"),
        # ORDER — ORDER_LINE — PRODUCT
        (7.28, 1.82, "(1,1)"),
        (6.18, 2.94, "(1,N)"),
        (5.42, 2.28, "(0,N)"),
        (5.58, 3.52, "(1,1)"),
        # ORDER — PROCESS — SALES_REP
        (9.32, 1.78, "(0,N)"),
        (11.44, 1.78, "(1,1)"),
        # PRODUCT — PRODUCE — WORK_CENTER
        (5.62, 4.16, "(0,N)"),
        (7.82, 3.62, "(1,1)"),
        # WORK_CENTER — ASSIGNMENT — EMPLOYEE
        (9.62, 3.62, "(1,1)"),
        (9.62, 4.38, "(5,N)"),
        (11.48, 3.62, "(1,1)"),
        (11.48, 4.38, "(1,N)"),
        # WORK_CENTER — MANAGE — MANAGER
        (9.78, 4.48, "(1,1)"),
        (9.78, 5.66, "(1,1)"),
    ]
    for x, y, text in labels:
        add_card_in_gap_xy(slide, x, y, text, obstacles)

    # leftover cards are everything after the original entity boxes
    entity_boxes = list(E.values())
    for ob in obstacles[len(entity_boxes):]:
        hits = [n for n, b in E.items() if _overlap(ob, b, pad=0.03)]
        if hits:
            print("LABEL OVERLAP", tuple(round(v, 2) for v in ob), hits)

    add_rect(slide, 0.18, 6.12, 0.38, 0.24, [""], size=8, line_pt=1.0)
    add_textbox(slide, 0.62, 6.12, 2.2, 0.24, "Entity", size=12)
    add_rect(slide, 0.18, 6.42, 0.38, 0.24, [""], size=8, line_pt=2.25)
    add_textbox(slide, 0.62, 6.42, 2.6, 0.24, "Composite entity (from M:N)", size=12)
    add_diamond(slide, 0.12, 6.70, 0.48, 0.32, [""], size=6)
    add_textbox(slide, 0.62, 6.74, 2.6, 0.24, "Relationship (1:1 or 1:N)", size=12)
    add_textbox(slide, 3.40, 6.74, 5.4, 0.24, "O  =  overlapping, partial specialization", size=12)


def add_schema_line(tf, first, rel, fields):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT
    p.space_before = Pt(8)
    p.space_after = Pt(2)
    p.line_spacing = 1.15
    r = p.add_run()
    r.text = rel + ":  "
    set_run_font(r, size=16, bold=False)
    for i, (attr, is_pk, is_fk) in enumerate(fields):
        if i:
            sep = p.add_run()
            sep.text = ",  "
            set_run_font(sep, size=16)
        ar = p.add_run()
        ar.text = attr
        set_run_font(ar, size=16, underline=is_pk)
        if is_fk:
            at = p.add_run()
            at.text = "@"
            set_run_font(at, size=16, underline=is_pk)


def build_schema(slide):
    add_textbox(
        slide, 0.40, 0.12, 12.5, 0.36,
        "95–703 C  Homework #2   —   Standard Answer:  Relational Schema",
        size=18,
    )
    add_textbox(
        slide, 0.40, 0.50, 12.5, 0.32,
        "PK underlined.   FK marked with @.   Identifying FK that is also part of the PK is both underlined and tagged @.",
        size=13, italic=True,
    )

    left = [
        ("Supplier", [
            ("Supplier_ID", True, False),
            ("Supplier_Name", False, False),
            ("Supplier_Address", False, False),
            ("Supplier_Phone", False, False),
        ]),
        ("Procurement", [
            ("Item_ID", True, True),
            ("Procurement_Date", True, False),
            ("Supplier_ID", False, True),
            ("Quantity", False, False),
            ("Unit_Price", False, False),
        ]),
        ("Item", [
            ("Item_ID", True, False),
            ("Description", False, False),
            ("Average_Unit_Cost", False, False),
            ("Units_On_Hand", False, False),
            ("Reorder_Point", False, False),
        ]),
        ("Product_Assembly", [
            ("Product_ID", True, True),
            ("Item_ID", True, True),
            ("Quantity_Needed", False, False),
        ]),
        ("Product", [
            ("Product_ID", True, False),
            ("Product_Name", False, False),
            ("Product_Line_ID", False, True),
            ("Work_Center_ID", False, True),
        ]),
        ("Product_Line", [
            ("Product_Line_ID", True, False),
            ("Product_Line_Name", False, False),
        ]),
        ("Customer", [
            ("Customer_ID", True, False),
            ("Customer_Name", False, False),
            ("Customer_Email", False, False),
            ("Customer_Address", False, False),
            ("Customer_Phone", False, False),
        ]),
        ("Order", [
            ("Order_ID", True, False),
            ("Order_Date", False, False),
            ("Customer_ID", False, True),
            ("Sales_Rep_ID", False, True),
        ]),
    ]
    right = [
        ("Order_Line", [
            ("Order_ID", True, True),
            ("Product_ID", True, True),
            ("Quantity_Ordered", False, False),
            ("Unit_Price", False, False),
        ]),
        ("Work_Center", [
            ("Work_Center_ID", True, False),
            ("Work_Center_Name", False, False),
            ("Manager_ID", False, True),
        ]),
        ("Assignment", [
            ("Employee_ID", True, True),
            ("Work_Center_ID", True, True),
        ]),
        ("Employee", [
            ("Employee_ID", True, False),
            ("First_Name", False, False),
            ("Last_Name", False, False),
            ("Address", False, False),
            ("Salary", False, False),
            ("Title", False, False),
        ]),
        ("Manager", [
            ("Employee_ID", True, True),
            ("Driver_License", False, False),
            ("Budget_Amount", False, False),
        ]),
        ("Admin_Assistant", [
            ("Employee_ID", True, True),
            ("Accounting_Training", False, False),
            ("MS_Office_Level", False, False),
        ]),
        ("Sales_Rep", [
            ("Employee_ID", True, True),
            ("Sales_Area", False, False),
            ("Email", False, False),
        ]),
    ]

    sp = slide.shapes.add_textbox(Inches(0.40), Inches(0.90), Inches(6.2), Inches(5.55))
    tf = _tf_setup(sp, anchor="t", wrap=True, l=0.04, t=0.02, r=0.04, b=0.02)
    for i, (rel, fields) in enumerate(left):
        add_schema_line(tf, i == 0, rel, fields)

    sp2 = slide.shapes.add_textbox(Inches(6.70), Inches(0.90), Inches(6.2), Inches(5.55))
    tf2 = _tf_setup(sp2, anchor="t", wrap=True, l=0.04, t=0.02, r=0.04, b=0.02)
    for i, (rel, fields) in enumerate(right):
        add_schema_line(tf2, i == 0, rel, fields)

    add_textbox(
        slide, 0.40, 6.55, 12.5, 0.55,
        "Notes:  Work_Center.Manager_ID@ is unique (1:1 with Manager).  "
        "Customer_Email may be null for potential customers.  "
        "Procurement PK = (Item_ID, Procurement_Date) so that at a given date an item has only one supplier.  "
        "Assignment: employee (1,N), work center (5,N).",
        size=12, italic=True,
    )


def main():
    shutil.copy(TEMPLATE, OUT_PPTX)
    prs = Presentation(OUT_PPTX)

    # Slide 1 of the template is the Painter ERD — replace with HW2 ER
    s1 = prs.slides[0]
    clear_slide(s1)
    build_er(s1)

    s2 = add_blank_slide(prs)
    build_schema(s2)

    # core props
    prs.core_properties.title = "95-703 C Homework #2 Standard Answer"
    prs.core_properties.subject = "ER diagram and relational schema"
    prs.core_properties.author = "TA Standard Answer"

    prs.save(OUT_PPTX)
    shutil.copy(OUT_PPTX, OUT_ROOT)
    desktop_dir = os.path.dirname(OUT_DESKTOP)
    os.makedirs(desktop_dir, exist_ok=True)
    shutil.copy(OUT_PPTX, OUT_DESKTOP)
    print("wrote", OUT_PPTX)
    print("wrote", OUT_ROOT)
    print("wrote", OUT_DESKTOP)


if __name__ == "__main__":
    main()
